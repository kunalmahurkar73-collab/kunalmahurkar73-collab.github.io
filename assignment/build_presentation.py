"""Build the evidence-first Amazon India checkout assignment deck."""
from pathlib import Path

from PIL import Image, ImageOps, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "assignment"
PPTX = OUT / "amazon-india-web-checkout-ux-evaluation-7-slides.pptx"
PDF = OUT / "amazon-india-web-checkout-ux-evaluation-7-slides.pdf"
PREVIEWS = OUT / "slide-previews"
NAVY = RGBColor(19, 25, 33)
INK = RGBColor(15, 17, 17)
MUTED = RGBColor(76, 87, 91)
TEAL = RGBColor(0, 113, 133)
AMBER = RGBColor(196, 85, 0)
PAPER = RGBColor(247, 248, 248)
WHITE = RGBColor(255, 255, 255)
LINE = RGBColor(213, 217, 217)

deck = Presentation()
deck.slide_width = Inches(13.333)
deck.slide_height = Inches(7.5)


def text(slide, value, x, y, w, h, size=18, color=INK, bold=False, align=None):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(.02)
    frame.margin_top = frame.margin_bottom = Inches(.01)
    for index, line in enumerate(value.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.font.name = "Arial"
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = color
        paragraph.space_after = Pt(7)
        if align is not None:
            paragraph.alignment = align
    return shape


def image(slide, path, x, y, w, h):
    path = Path(path)
    with Image.open(path) as original:
        ratio = original.width / original.height
    width = min(w, h * ratio)
    height = min(h, w / ratio)
    slide.shapes.add_picture(str(path), Inches(x + (w - width) / 2), Inches(y + (h - height) / 2), width=Inches(width), height=Inches(height))


def evidence(name):
    cleaned = ROOT / "evidence" / "clean" / Path(name).with_suffix(".png")
    return cleaned if cleaned.exists() else ROOT / "evidence" / name


def design(number):
    return OUT / "design-previews" / f"screen-{number:02d}.png"


def slide(title, number, note, subtitle=None):
    item = deck.slides.add_slide(deck.slide_layouts[6])
    number = len(deck.slides)
    item.background.fill.solid()
    item.background.fill.fore_color.rgb = PAPER
    text(item, f"{number:02d}  AMAZON INDIA CHECKOUT", .62, .27, 5.4, .28, 10, TEAL, True)
    text(item, title, .62, .70, 12.05, .68, 29, NAVY, True)
    if subtitle:
        text(item, subtitle, .64, 1.40, 12.0, .48, 15, MUTED)
    text(item, "Kunal Mahurkar · Independent UX evaluation · 20 Oct 2026", .62, 7.16, 10.5, .18, 9, MUTED)
    text(item, f"{number:02d} / 07", 11.92, 7.15, .77, .18, 9, MUTED, align=PP_ALIGN.RIGHT)
    item.notes_slide.notes_text_frame.text = note
    return item


def label(item, value, x, y, w, color=TEAL):
    text(item, value.upper(), x, y, w, .28, 11, color, True)


def point(item, heading, body, x, y, w, body_h=.75):
    text(item, heading, x, y, w, .36, 17, NAVY, True)
    text(item, body, x, y + .43, w, body_h, 15, MUTED)


s = slide("What the checkout shows, and what I changed", 1,
          "I observed one Amazon India desktop checkout, placed a cash-on-delivery order and cancelled it. This deck separates visible facts from usability hypotheses and shows the redesign response.",
          "One checkout journey ? six heuristic concerns ? three UX laws")
text(s, "REAL INTERFACE", .75, 2.08, 3.7, .32, 13, TEAL, True)
text(s, "Seven public screenshot crops from one checkout and cancellation", .75, 2.52, 4.25, 1.12, 23, NAVY, True)
text(s, "REDESIGN", .75, 4.25, 3.7, .32, 13, AMBER, True)
text(s, "Editable screen designs and a working checkout simulation", .75, 4.68, 4.25, 1.12, 23, NAVY, True)
image(s, design(5), 5.55, 1.99, 6.87, 3.87)
text(s, "Proposed review screen. This is a concept, not an Amazon screen.", 5.6, 6.03, 6.72, .42, 12, MUTED)

s = slide("Evidence: one observed journey", 2,
          "These are genuine public crops from one desktop browser journey. They establish the displayed UI, not shopper confusion or conversion performance.",
          "Seven original crops. One cash-on-delivery order placed and cancelled immediately.")
image(s, evidence("address-stage-header.jpg"), .70, 2.05, 3.76, 1.45)
image(s, evidence("delivery-options.jpg"), 4.78, 2.03, 3.76, 2.42)
image(s, evidence("cancellation-success.jpg"), 8.73, 2.03, 3.82, 2.42)
label(s, "EV-02 · Address", .77, 4.67, 3.5)
label(s, "EV-04 · Delivery", 4.85, 4.67, 3.4)
label(s, "EV-07 · Cancellation", 8.8, 4.67, 3.4)
text(s, "Address choices are indicated; personal rows are withheld.", .77, 5.07, 3.45, .92, 15)
text(s, "Two delivery choices appear, both marked free.", 4.85, 5.07, 3.4, .92, 15)
text(s, "Cancellation completed after the observed order.", 8.8, 5.07, 3.4, .92, 15)
text(s, "Evidence boundary: no participants, analytics baseline, or real payment failure were observed.", .77, 6.37, 11.75, .43, 15, AMBER, True)

s = slide("01  Final cost may be easy to miss", 3,
          "Observed: item 207 rupees, cash-on-delivery fee 7 rupees, total 214 rupees. Visibility of system status is the relevant heuristic. The fee is visible, so missed comprehension is a hypothesis, not a proven violation. Test by asking for the total and fee unaided.",
          "Nielsen #1: visibility of system status · Current fact versus design response")
label(s, "Observed on Amazon · EV-05", .75, 2.02, 5.48)
image(s, evidence("review-total.jpg"), .72, 2.45, 5.68, 2.10)
point(s, "₹207 item + ₹7 COD fee = ₹214 total", "The final amount is visible at review. Whether shoppers notice the fee before commitment is unmeasured.", .77, 4.83, 5.45, 1.0)
label(s, "Proposed review", 6.87, 2.02, 5.55, AMBER)
image(s, design(5), 6.84, 2.45, 5.63, 2.95)
text(s, "Itemized summary beside the decision; explicit total payable and Edit actions.", 6.9, 5.65, 5.44, .73, 16)
text(s, "Test: ask shoppers to state the total and the reason for the extra ₹7 before clicking.", .78, 6.65, 11.75, .31, 14, TEAL, True)

s = slide("02  Payment methods need cost context", 4,
          "The screenshot shows cards, net banking, UPI scan and pay, and cash on delivery. The possible comprehension cost is a hypothesis. Nielsen's match with the real world and Hick's Law guide simpler labels and fee disclosure. No selection time was measured.",
          "Nielsen #2: match to familiar language · Hick's Law: decision effort")
label(s, "Observed on Amazon · EV-03", .74, 1.94, 5.52)
image(s, evidence("payment-methods.jpg"), .73, 2.30, 5.45, 3.50)
label(s, "Proposed payment design", 6.79, 1.94, 5.63, AMBER)
image(s, design(4), 6.70, 2.30, 5.72, 3.35)
text(s, "Visible fact: multiple payment labels and an insufficient balance message.", .75, 5.98, 5.57, .57, 15)
text(s, "Change: group methods, use plain labels, and show the COD fee at selection.", 6.82, 5.98, 5.53, .57, 15)
text(s, "Test: compare method comprehension and selection time; a current violation is unproven.", .75, 6.68, 11.68, .31, 14, TEAL, True)

s = slide("03  Choices and controls compete for attention", 5,
          "The cart crop shows item controls and promotion content. This can raise choice and pointer concerns, but a screenshot does not measure hesitation or missed clicks. Nielsen's minimalist design and flexibility heuristics, plus Hick's and Fitts's laws, motivate grouping and more consistent click areas.",
          "Nielsen #8 and #7 · Hick's Law and Fitts's Law")
label(s, "Observed cart · EV-01", .75, 1.95, 5.52)
image(s, evidence("cart-item.jpg"), .73, 2.31, 5.68, 2.85)
label(s, "Proposed cart", 6.85, 1.95, 5.49, AMBER)
image(s, design(1), 6.78, 2.31, 5.62, 2.85)
point(s, "Hypothesis: competing decisions", "Promotion content and quantity control share the cart view. Decision effort has not been timed.", .78, 5.45, 5.48, .75)
point(s, "Change: clearer action hierarchy", "Keep quantity and continue actions prominent, with larger consistent click areas.", 6.89, 5.45, 5.46, .75)

s = slide("Heuristic map: observation, risk, response", 8,
          "The six rows summarize which Nielsen heuristic is relevant. None is presented as a verified user harm. Five observations come from the real journey; payment failure is a proposed scenario. Severity is expert judgement and must be checked through user testing.",
          "A heuristic concern is an expert assessment. A user impact requires a user study.")
rows = [
    ("#1 Visibility", "₹7 COD fee and ₹214 total visible", "Fee comprehension unknown", "Persistent itemized total"),
    ("#2 Real-world match", "Several payment labels shown", "Meaning or fee may need context", "Plain labels and fee at choice"),
    ("#3 User control", "Change links observed at review", "Rework not established", "Direct Edit, retained choices"),
    ("#7 Efficiency", "Quantity and Change controls visible", "Pointer misses unmeasured", "Consistent click targets"),
    ("#8 Minimalism", "Promotions and choices share view", "Decision effort unmeasured", "Clearer hierarchy"),
    ("#9 Error recovery", "No failure captured", "Scenario only", "Explicit status and safe retry"),
]
table = s.shapes.add_table(7, 4, Inches(.66), Inches(2.06), Inches(12.0), Inches(4.08)).table
for i, width in enumerate((2.11, 3.16, 3.16, 3.57)):
    table.columns[i].width = Inches(width)
for column, heading in enumerate(("Heuristic", "Observed", "Risk or limit", "Proposed change")):
    table.cell(0, column).text = heading
for row, values in enumerate(rows, 1):
    for column, value in enumerate(values):
        table.cell(row, column).text = value
for row in range(7):
    for col in range(4):
        cell = table.cell(row, col)
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY if row == 0 else (WHITE if row % 2 else PAPER)
        cell.margin_left = Inches(.13)
        cell.margin_right = Inches(.09)
        cell.margin_top = Inches(.10)
        cell.margin_bottom = Inches(.06)
        for paragraph in cell.text_frame.paragraphs:
            paragraph.font.name = "Arial"
            paragraph.font.size = Pt(12 if row else 13)
            paragraph.font.bold = row == 0 or col == 0
            paragraph.font.color.rgb = WHITE if row == 0 else INK
text(s, "Hick: group choices ? Fitts: enlarge targets ? Jakob: familiar edits. All impacts await testing.", .70, 6.42, 11.85, .34, 14, AMBER, True)

s = slide("Next step: test the proposed improvements", 10,
          "Priorities: cost comprehension and risk-led recovery first, then payment choice and editing, then target polish. Recruit five India-based desktop shoppers. Ask for the total and fee, a late edit, and recovery from a simulated failure. Measure unaided task completion, time, errors, missed clicks and confidence. No outcomes are claimed yet. Sources: original screenshots; Nielsen Norman Group heuristics; Interaction Design Foundation summaries of Hick's and Fitts's laws.",
          "Priority reflects expert judgement. No usability sessions or outcome metrics have been measured.")
point(s, "P1 · Total and fee", "Ask shoppers to state the amount payable and explain the COD fee before placing an order.", .77, 2.03, 5.51, .84)
point(s, "P1 discovery · Failure recovery", "Use a controlled simulated failure. Check whether shoppers understand payment status and recover safely.", 6.82, 2.03, 5.44, .84)
point(s, "P2 · Payment and editing", "Compare method comprehension and late-edit success without prompting.", .77, 3.64, 5.51, .80)
point(s, "P3 · Pointer targets", "Measure target sizes and missed clicks during desktop tasks.", 6.82, 3.64, 5.44, .80)
text(s, "Measures: correct total and fee, unaided task success, time, errors, recovery success, confidence.", .77, 5.38, 11.67, .59, 17, NAVY, True)
text(s, "References: nngroup.com/articles/ten-usability-heuristics/ · interaction-design.org/literature/topics/hick-s-law · interaction-design.org/literature/topics/fitts-law", .77, 6.18, 11.7, .62, 10, MUTED)

OUT.mkdir(exist_ok=True)
deck.save(PPTX)
print(f"Saved {PPTX} ({len(deck.slides)} slides)")

import win32com.client
app = win32com.client.DispatchEx("PowerPoint.Application")
presentation = app.Presentations.Open(str(PPTX), WithWindow=False)
try:
    presentation.SaveAs(str(PDF), 32)
    PREVIEWS.mkdir(exist_ok=True)
    thumbs = []
    for index, item in enumerate(presentation.Slides, 1):
        path = PREVIEWS / f"slide-{index:02d}.png"
        item.Export(str(path), "PNG", 1280, 720)
        with Image.open(path) as opened:
            thumbs.append(opened.convert("RGB").resize((512, 288)))
finally:
    presentation.Close()
    app.Quit()

sheet = Image.new("RGB", (512 * 2, 288 * 4), "#ffffff")
for index, thumb in enumerate(thumbs):
    sheet.paste(thumb, ((index % 2) * 512, (index // 2) * 288))
sheet.save(PREVIEWS / "contact-sheet.png")
for stale in PREVIEWS.glob("slide-*.png"):
    if int(stale.stem.rsplit("-", 1)[1]) > len(thumbs):
        stale.unlink()
print(f"Saved {PDF} and {len(thumbs)} slide previews")
