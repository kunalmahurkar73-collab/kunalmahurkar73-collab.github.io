"""Export native PowerPoint shapes and text as Penpot-importable editable SVGs."""
from pathlib import Path
from html import escape
from pptx import Presentation

HERE=Path(__file__).resolve().parent
SOURCE=HERE/'amazon-checkout-editable-design-source.pptx'
OUT=HERE/'penpot-import'
OUT.mkdir(exist_ok=True)
deck=Presentation(SOURCE)
EMU_PER_PX=9525

def px(value): return round(value/EMU_PER_PX,2)
def rgb(color, fallback='none'):
    try: return f'#{color.rgb}' if color and color.rgb else fallback
    except (AttributeError, TypeError): return fallback

for number,slide in enumerate(deck.slides,1):
    lines=['<svg xmlns="http://www.w3.org/2000/svg" width="1280" height="720" viewBox="0 0 1280 720">']
    lines.append(f'<title>Amazon checkout redesign screen {number:02d}</title>')
    lines.append('<rect width="1280" height="720" fill="#f5f6f6"/>')
    for shape in slide.shapes:
        x,y,w,h=map(px,(shape.left,shape.top,shape.width,shape.height))
        if shape.has_text_frame and shape.text.strip():
            for i,p in enumerate(shape.text_frame.paragraphs):
                if not p.text: continue
                f=p.font
                size=round(f.size.pt*96/72,2) if f.size else 18
                weight='700' if f.bold else '400'
                lines.append(f'<text x="{x+3}" y="{round(y+size*(i+1)+2,2)}" fill="{rgb(f.color,"#131921")}" font-family="Arial, sans-serif" font-size="{size}" font-weight="{weight}">{escape(p.text)}</text>')
        else:
            fill=rgb(shape.fill.fore_color,'#ffffff')
            try: stroke=rgb(shape.line.color,'none')
            except (AttributeError,TypeError): stroke='none'
            radius=' rx="10"' if shape.auto_shape_type else ''
            lines.append(f'<rect x="{x}" y="{y}" width="{w}" height="{h}"{radius} fill="{fill}" stroke="{stroke}"/>')
    lines.append('</svg>')
    (OUT/f'screen-{number:02d}.svg').write_text('\n'.join(lines),encoding='utf-8')
print(f'{len(deck.slides)} SVG screens: {OUT}')
